from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def export_pptx(document: dict[str, Any], target: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for slide_data in document.get("slides", []):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _rgb("F6F8FB")
        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.7))
        title_frame = title_box.text_frame
        title_frame.clear()
        paragraph = title_frame.paragraphs[0]
        paragraph.text = str(slide_data.get("title", ""))
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(24)
        paragraph.font.bold = True
        paragraph.font.color.rgb = _rgb("152238")
        for element in slide_data.get("elements", []):
            _add_element(slide, element)
        if slide_data.get("notes"):
            notes = slide.notes_slide.notes_text_frame
            notes.text = str(slide_data["notes"])
    if not presentation.slides:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    target.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(target))


def _add_element(slide: Any, element: dict[str, Any]) -> None:
    x = Inches(float(element.get("x", 0.6)))
    y = Inches(float(element.get("y", 1.2)))
    w = Inches(float(element.get("w", 12.0)))
    h = Inches(float(element.get("h", 5.5)))
    kind = element.get("type", "text")
    if kind == "image" and element.get("src"):
        image = Path(str(element["src"]))
        if image.exists():
            slide.shapes.add_picture(str(image), x, y, width=w, height=h)
            return
    if kind == "table":
        rows = element.get("rows", [])
        if rows:
            columns = max(len(row) for row in rows)
            table = slide.shapes.add_table(len(rows), columns, x, y, w, h).table
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = str(value)
            return
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    text = str(element.get("text", element.get("content", "")))
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Microsoft YaHei"
    paragraph.font.size = Pt(float(element.get("font_size", 16)))
    paragraph.font.color.rgb = _rgb(str(element.get("color", "26364D")).lstrip("#"))
    paragraph.alignment = PP_ALIGN.LEFT


def _rgb(value: str) -> Any:
    from pptx.dml.color import RGBColor

    normalized = value[:6].ljust(6, "0")
    return RGBColor(  # type: ignore[no-untyped-call]
        int(normalized[0:2], 16),
        int(normalized[2:4], 16),
        int(normalized[4:6], 16),
    )


def default_report(run_id: str, project: dict[str, Any], data: dict[str, Any]) -> dict[str, Any]:
    sections = ["项目说明", "计算模型", "网格信息", "材料与属性", "边界条件", "工况说明", "应力结果", "位移结果", "热点详情", "结论", "附录"]
    slides: list[dict[str, Any]] = [
        {
            "id": "cover",
            "title": project["name"],
            "layout": "cover",
            "elements": [
                {"type": "text", "x": 0.8, "y": 2.5, "w": 11.7, "h": 1.2, "font_size": 28, "text": "CAE 结果分析报告"},
                {"type": "text", "x": 0.8, "y": 4.0, "w": 11.7, "h": 0.5, "text": f"项目编号：{project.get('code') or '-'}"},
            ],
            "notes": "",
        }
    ]
    for index, section in enumerate(sections, start=1):
        elements: list[dict[str, Any]] = [
            {"type": "text", "x": 0.7, "y": 1.3, "w": 11.8, "h": 4.8, "text": _section_text(section, project, data)}
        ]
        slides.append({"id": f"slide-{index}", "title": section, "layout": "title-content", "elements": elements, "notes": ""})
    return {
        "id": "",
        "run_id": run_id,
        "template_id": "default-v1",
        "page_size": "LAYOUT_WIDE",
        "theme": {"primary": "#176B87", "font": "Microsoft YaHei"},
        "sections": [{"id": f"section-{i}", "title": title} for i, title in enumerate(sections)],
        "slides": slides,
    }


def _section_text(section: str, project: dict[str, Any], data: dict[str, Any]) -> str:
    if section == "项目说明":
        return project.get("description") or f"{project['name']} / {project.get('analysis_type', '')}"
    if section == "计算模型":
        return json.dumps(data.get("model_summary", {}), ensure_ascii=False, indent=2)
    if section == "网格信息":
        return json.dumps(data.get("mesh_summary", {}), ensure_ascii=False, indent=2)
    if section == "材料与属性":
        return f"材料 {len(data.get('materials', []))} 项；属性 {len(data.get('properties', []))} 项。"
    if section == "边界条件":
        boundary = data.get("boundary_conditions", {})
        return f"载荷 {len(boundary.get('loads', []))}；约束 {len(boundary.get('constraints', []))}；连接 {len(boundary.get('connections', []))}。"
    if section == "工况说明":
        return f"共识别 {len(data.get('subcases', []))} 个子工况。"
    if section in {"应力结果", "位移结果", "热点详情"}:
        hotspots = data.get("hotspots", [])
        return f"已识别 {len(hotspots)} 个结果热点，详细图片和审核结论见工作台。"
    if section == "结论":
        return "本页由工程师填写。系统不自动判定工程是否合格。"
    return "分析输入、软件版本、任务状态和诊断信息见项目工作台。"
